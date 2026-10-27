import json, re, time
from datetime import datetime, timezone
from pathlib import Path
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
HEADERS = ["Discovered Date", "Job Title", "Company", "Location", "Work Mode", "Job Category", "Experience Required", "Fresher Eligible", "Required Skills", "Matched Skills", "Match Score (%)", "Salary", "Posted Date", "Deadline", "Job URL", "Source", "Listing Status", "Application Status", "Notes"]
DEFAULTS = {
    "output_file": "job_tracker.xlsx", "minimum_match_score": 20,
    "pages_per_source": 2, "results_per_page": 20,
    "himalayas_search_pages": 2,
    "role_keywords": ["soc analyst", "soc", "cybersecurity analyst", "cyber security analyst", "security analyst", "information security", "vapt", "penetration tester", "vulnerability assessment", "security operations", "network support", "network engineer", "noc engineer", "network security", "it support", "technical support", "service desk", "help desk", "system administrator", "cybersecurity intern", "security intern", "networking intern", "it intern", "internship", "graduate", "fresher", "entry level", "junior security analyst"],
    "skill_keywords": ["python", "linux", "networking", "tcp/ip", "dns", "dhcp", "http", "https", "nmap", "wireshark", "tcpdump", "burp suite", "vulnerability assessment", "penetration testing", "incident response", "siem", "flask", "sql", "windows", "firewall", "network security", "soc", "splunk", "security operations"],
    "fresher_signals": ["fresher", "entry level", "entry-level", "junior", "intern", "internship", "graduate", "new grad", "0-1 year", "0 - œ1 year", "no experience", "trainee", "early career", "associate"],
    "search_queries": ["cybersecurity", "SOC analyst", "security analyst", "VAPT", "network support", "NOC engineer", "IT support", "security internship"]
}
S = requests.Session()
S.headers.update({"User-Agent": "PersonalJobSearchAgent/1.1 (personal job discovery)", "Accept": "application/json"})


def cfg_load():
    c = DEFAULTS.copy()
    p = HERE / "config.json"
    if p.exists():
        try:
            c.update(json.loads(p.read_text(encoding="utf-8-sig")))
        except Exception as e:
            print("Warning reading config.json; using defaults:", e)
    return c


def txt(v):
    if v is None: return ""
    if isinstance(v, list): return ", ".join(map(str, v))
    if isinstance(v, dict): return json.dumps(v, ensure_ascii=False)
    return re.sub(r"\s+", " ", str(v)).strip()


def get(url, params=None):
    r = S.get(url, params=params, timeout=30)
    r.raise_for_status()
    return r.json()


def date_value(v):
    if isinstance(v, (int, float)):
        try: return datetime.fromtimestamp(v, tz=timezone.utc).strftime("%Y-%m-%d")
        except (ValueError, OSError, OverflowError): return str(v)
    return txt(v)


def himalayas(c):
    """Search Himalayas' public API for remote jobs open to India-based applicants."""
    out, seen = [], set()
    pages = max(1, min(int(c.get("himalayas_search_pages", 2)), 5))
    queries = c.get("search_queries") or DEFAULTS["search_queries"]
    for q in queries:
        for page in range(1, pages + 1):
            try:
                d = get("https://himalayas.app/jobs/api/search", {"q": q, "country": "IN", "sort": "recent", "page": page})
            except Exception as e:
                print(f"    Himalayas query '{q}' page {page} failed: {e}")
                break
            batch = d.get("jobs", []) if isinstance(d, dict) else []
            if not batch: break
            for x in batch:
                url = txt(x.get("applicationLink"))  # API terms specify applicationLink for Apply actions.
                guid = txt(x.get("guid")) or url
                if not url or guid in seen: continue
                seen.add(guid)
                restrictions = x.get("locationRestrictions") or []
                countries = [txt(r.get("name") if isinstance(r, dict) else r) for r in restrictions]
                loc = "Remote  -   " + (", ".join(countries) if countries else "Worldwide / no country restriction")
                country_ok = not restrictions or any(isinstance(r, dict) and (str(r.get("alpha2", "")).upper() == "IN" or "india" in str(r.get("name", "")).lower()) for r in restrictions)
                note = "Himalayas country search matched India; verify employer requirements before applying."
                if not country_ok: note = "Check the source country restrictions carefully before applying."
                salary = ""
                if x.get("minSalary") is not None or x.get("maxSalary") is not None:
                    salary = f"{txt(x.get('currency'))} {x.get('minSalary', '?')} - œ{x.get('maxSalary', '?')} {txt(x.get('salaryPeriod'))}".strip()
                out.append({"title": x.get("title"), "company": x.get("companyName") or x.get("company"), "location": loc, "mode": "Remote", "description": x.get("description") or x.get("excerpt"), "skills": x.get("categories") or x.get("parentCategories") or [], "salary": salary, "posted": date_value(x.get("pubDate")), "url": url, "source": "Himalayas", "experience": txt(x.get("seniority")) or txt(x.get("employmentType")), "eligibility_note": note})
            if len(batch) < 20: break
            time.sleep(0.25)
    return out


def remoteok(c):
    d = get("https://remoteok.com/api")
    out = []
    if not isinstance(d, list): return out
    for x in d:
        if not isinstance(x, dict) or not x.get("position"): continue
        jid = x.get("id")
        out.append({"title": x.get("position"), "company": x.get("company"), "location": x.get("location") or "Remote  -   India eligibility unclear", "mode": "Remote", "description": x.get("description"), "skills": x.get("tags") or [], "salary": " - ".join(str(x[k]) for k in ("salary_min", "salary_max") if x.get(k)), "posted": x.get("date"), "url": x.get("url") or (f"https://remoteok.com/remote-jobs/{jid}" if jid else ""), "source": "Remote OK", "experience": "", "eligibility_note": "This feed does not confirm India eligibility; check the original listing."})
    return out


def arbeitnow(c):
    out = []
    for page in range(1, max(1, min(int(c["pages_per_source"]), 5)) + 1):
        d = get("https://www.arbeitnow.com/api/job-board-api", {"page": page})
        batch = d.get("data", []) if isinstance(d, dict) else []
        if not batch: break
        for x in batch:
            rem = bool(x.get("remote"))
            out.append({"title": x.get("title"), "company": x.get("company_name"), "location": x.get("location") or ("Remote  -   India eligibility unclear" if rem else "Location not specified"), "mode": "Remote" if rem else "On-site/Hybrid  -   verify", "description": x.get("description"), "skills": x.get("tags") or [], "salary": "", "posted": date_value(x.get("created_at")), "url": x.get("url"), "source": "Arbeitnow", "experience": "", "eligibility_note": "Check the employer listing for India location and work eligibility."})
        time.sleep(0.25)
    return out


def category(title, desc):
    t = (txt(title) + " " + txt(desc)).lower()
    if any(k in t for k in ["cybersecurity", "cyber security", "soc analyst", "vapt", "penetration tester", "vulnerability", "information security", "security analyst", "security operations"]): return "Cybersecurity"
    if any(k in t for k in ["network", "noc engineer"]): return "Networking"
    if any(k in t for k in ["it support", "technical support", "service desk", "help desk", "system administrator"]): return "IT Support"
    return "Other / Review"


def score(x, c):
    title, desc, skills = txt(x.get("title")), txt(x.get("description")), txt(x.get("skills"))
    combined = f"{title} {desc} {skills}".lower()
    matched = [k for k in c["skill_keywords"] if k.lower() in combined]
    roles = [k for k in c["role_keywords"] if k.lower() in title.lower()]
    points = min(100, round(len(matched) * 8 + min(len(roles), 2) * 15))
    fresh = any(k.lower() in combined for k in c["fresher_signals"])
    if any(k in combined for k in ["5+ years", "5 years", "senior", "principal", "staff engineer", "lead engineer", "10+ years"]):
        fresh = False
        points = max(0, points - 20)
    return points, matched, fresh



def location_relevance(job):
    """Label location evidence without claiming eligibility the source does not prove."""
    location = txt(job.get("location")).lower()
    source = txt(job.get("source")).lower()
    note = txt(job.get("eligibility_note"))
    india_terms = ("india", "bengaluru", "bangalore", "hyderabad", "chennai", "pune",
                   "mumbai", "new delhi", "gurugram", "gurgaon", "noida", "kolkata",
                   "ahmedabad", "kochi", "thiruvananthapuram")
    if any(term in location for term in india_terms):
        return "India location indicated", note or "Confirm the exact work location and hiring eligibility on the employer page."
    if source == "himalayas" and ("india" in location or "worldwide / no country restriction" in location.lower()):
        return "Remote  -   India search / unrestricted; verify", note or "Confirm India eligibility, timezone, and work authorization on the employer page."
    if "remote" in location or txt(job.get("mode")).lower() == "remote":
        return "Remote  -   India eligibility unconfirmed", note or "Remote does not necessarily mean worldwide; confirm India is an eligible hiring location."
    if location and location not in ("location not specified", "unknown"):
        return "Location outside/unclear for India  -   verify", note or "Check whether the employer accepts applicants based in India before applying."
    return "Location unclear  -   verify", note or "Location is unclear in the feed; check the original employer listing."


def experience_assessment(job, fresher_signal):
    """Estimate experience fit conservatively; never reject unclear requirements."""
    text = (txt(job.get("title")) + " " + txt(job.get("description")) + " " + txt(job.get("experience"))).lower()
    # Clear early-career signals first.
    if re.search(r"\b(0\s*[- - œ]\s*2|0\s*to\s*2|0\s*[- - œ]\s*1|1\s*[- - œ]\s*2)\s*(years?|yrs?)\b", text):
        return "Likely suitable  -   0 - œ2 years indicated"
    if any(x in text for x in ["no experience", "fresher", "entry-level", "entry level", "new grad", "graduate role", "graduate program", "graduate programme", "internship", "intern", "trainee", "early career"]):
        return "Potential fresher fit  -   verify requirements"
    if re.search(r"\b(3|4|5|6|7|8|9|10)\s*\+?\s*(years?|yrs?)\b", text) or any(x in text for x in ["senior", "principal", "staff engineer", "lead engineer", "manager"]):
        return "May require 3+ years / experienced  -   review manually"
    if fresher_signal:
        return "Potential entry-level fit  -   verify requirements"
    return "Experience unclear  -   review manually"


def posted_age_days(value):
    """Return age in days when the feed provides a parseable date; otherwise None."""
    raw = txt(value)
    if not raw:
        return None
    try:
        if raw.isdigit():
            number = int(raw)
            # Treat large numbers as Unix timestamps.
            if number > 10_000_000:
                posted = datetime.fromtimestamp(number, tz=timezone.utc).date()
            else:
                return None
        else:
            candidate = raw.replace("Z", "+00:00")
            try:
                parsed = datetime.fromisoformat(candidate)
            except ValueError:
                parsed = None
                for fmt, candidate_value in (
                    ("%Y-%m-%d", raw[:10]),
                    ("%d-%m-%Y", raw[:10]),
                    ("%Y/%m/%d", raw[:10]),
                    ("%b %d, %Y", raw[:12]),
                    ("%B %d, %Y", raw[:20]),
                ):
                    try:
                        parsed = datetime.strptime(candidate_value.strip(), fmt)
                        break
                    except ValueError:
                        continue
                if parsed is None:
                    return None
            posted = parsed.date()
        return (datetime.now(timezone.utc).date() - posted).days
    except (ValueError, TypeError, OSError, OverflowError):
        return None


def freshness_label(value):
    age = posted_age_days(value)
    if age is None:
        return "Posting date unclear  -   verify"
    if age < 0:
        return "Future/invalid date  -   verify"
    if age <= 7:
        return f"Recent  -   {age} day(s) old"
    if age <= 30:
        return f"Posted {age} days ago  -   verify open"
    return f"Potentially stale  -   {age} days old; verify"


def style_tracker_sheet(ws):
    """Add filters and practical formatting while preserving existing tracker columns/data."""
    if ws.max_row == 1 and all(ws.cell(1, i).value is None for i in range(1, ws.max_column + 1)):
        ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([15, 32, 24, 38, 22, 20, 36, 32, 32, 32, 16, 22, 22, 18, 48, 16, 28, 20, 65], 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def open_sheet(path):
    import shutil

    wb = load_workbook(path) if path.exists() else Workbook()

    if "Agent Results" not in wb.sheetnames:
        ws = wb.create_sheet("Agent Results")
        ws.append(HEADERS)
    else:
        ws = wb["Agent Results"]
        old_headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]

        if old_headers[:len(HEADERS)] != HEADERS:
            backup = path.with_name(
                f"{path.stem}_before_schema_change_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            )
            if path.exists():
                shutil.copy2(path, backup)

            # Preserve all existing columns, including any custom columns.
            extra_headers = [
                h for h in old_headers
                if h is not None and h not in HEADERS
            ]
            target_headers = HEADERS + extra_headers

            old_ws = ws
            new_ws = wb.create_sheet("Agent Results Migration")
            new_ws.append(target_headers)

            old_header_positions = {}
            for col, header in enumerate(old_headers, start=1):
                if header is not None and header not in old_header_positions:
                    old_header_positions[header] = col

            for row_num in range(2, old_ws.max_row + 1):
                values = []
                for header in target_headers:
                    old_col = old_header_positions.get(header)
                    values.append(
                        old_ws.cell(row_num, old_col).value if old_col else None
                    )
                new_ws.append(values)

            wb.remove(old_ws)
            new_ws.title = "Agent Results"
            ws = new_ws

    style_tracker_sheet(ws)

    if "Job Tracker" in wb.sheetnames:
        style_tracker_sheet(wb["Job Tracker"])

    return wb, ws


def main():
    c = cfg_load()
    path = HERE / c["output_file"]
    alljobs = []

    print("AI Job Search Agent  -   India-focused update")
    print("Existing jobs will be reassessed; application statuses and personal notes will be preserved.")

    for name, fn in [
        ("Himalayas (India search)", himalayas),
        ("Remote OK", remoteok),
        ("Arbeitnow", arbeitnow),
    ]:
        try:
            jobs = fn(c)
            alljobs.extend(jobs)
            print(f"  {name}: retrieved {len(jobs)} listings")
        except Exception as e:
            print(f"  {name}: unavailable or response error: {e}")

    wb, ws = open_sheet(path)

    # Use header names rather than assuming column positions.
    col = {
        str(ws.cell(1, i).value): i
        for i in range(1, ws.max_column + 1)
        if ws.cell(1, i).value is not None
    }

    if "Job URL" not in col:
        wb.close()
        raise RuntimeError("The Agent Results worksheet has no Job URL column.")

    known = {}
    for row_num in range(2, ws.max_row + 1):
        value = ws.cell(row_num, col["Job URL"]).value
        if value:
            known[str(value).strip()] = row_num

    seen_urls = set()
    added = updated = skipped = 0
    today = datetime.now().strftime("%Y-%m-%d")
    minimum_score = int(c["minimum_match_score"])

    def get_value(row_num, header):
        if header not in col:
            return None
        return ws.cell(row_num, col[header]).value

    def set_value(row_num, header, value):
        if header in col:
            ws.cell(row_num, col[header]).value = value

    def nonempty(value):
        return value is not None and str(value).strip() != ""

    for x in alljobs:
        title = txt(x.get("title"))
        url = txt(x.get("url")).strip()

        if not title or not url:
            skipped += 1
            continue

        # Prevent duplicate processing when sources overlap.
        if url in seen_urls:
            skipped += 1
            continue
        seen_urls.add(url)

        points, matched, fresh = score(x, c)
        cat = category(title, x.get("description"))
        row_num = known.get(url)
        is_existing = row_num is not None

        # New jobs must pass the configured filters. Existing jobs remain in
        # the workbook even if their score changes or they no longer qualify.
        if not is_existing and (
            points < minimum_score or cat == "Other / Review"
        ):
            skipped += 1
            continue

        location_status, location_note = location_relevance(x)
        experience_status = experience_assessment(x, fresh)
        freshness = freshness_label(x.get("posted"))

        experience = txt(x.get("experience")) or "Not specified  -   verify"
        skills = x.get("skills") or []
        if isinstance(skills, (list, tuple, set)):
            skills_text = ", ".join(str(v) for v in skills if v)
        else:
            skills_text = txt(skills)

        matched_text = ", ".join(matched)
        company = txt(x.get("company"))
        location = txt(x.get("location"))
        mode = txt(x.get("mode"))
        salary = txt(x.get("salary"))
        posted = txt(x.get("posted"))
        source_name = txt(x.get("source"))

        eligibility_note = txt(x.get("eligibility_note"))
        assessment = (
            f"Agent assessment: India relevance: {location_status}. "
            f"Posting freshness: {freshness}. "
            "Verify that the vacancy is still open, the employer is legitimate, "
            "and all location, work authorization, education, and experience "
            "requirements before applying. "
            f"{location_note} {eligibility_note}"
        ).strip()

        listing_status = (
            "Potentially stale  -   verify"
            if freshness.startswith("Potentially stale")
            else "Needs verification"
        )

        if is_existing:
            # Refresh information from the latest source data when available.
            # Do not change Discovered Date, Deadline, Job URL, Application
            # Status, or user-entered personal notes.
            updates = {
                "Job Title": title,
                "Company": company,
                "Location": location,
                "Work Mode": mode,
                "Job Category": cat,
                "Experience Required": experience,
                "Fresher Eligible": experience_status,
                "Required Skills": skills_text,
                "Matched Skills": matched_text,
                "Match Score (%)": points,
                "Salary": salary,
                "Posted Date": posted,
                "Source": source_name,
                "Listing Status": listing_status,
            }

            for header, value in updates.items():
                # Do not erase existing information with an empty source field.
                if header == "Job Category" and value == "Other / Review":
                    continue
                if header == "Experience Required" and value == "Not specified  -   verify":
                    if nonempty(get_value(row_num, header)):
                        continue
                if nonempty(value):
                    set_value(row_num, header, value)

            # Replace only the agent-managed assessment section. Everything
            # before it is retained as existing notes.
            old_notes = txt(get_value(row_num, "Notes"))
            marker = "Agent assessment:"
            personal_notes = old_notes.split(marker, 1)[0].rstrip()
            combined_notes = (
                personal_notes + "\n\n" + assessment
                if personal_notes
                else assessment
            )
            set_value(row_num, "Notes", combined_notes)
            updated += 1

        else:
            # Build new rows by header name, keeping the expected schema intact.
            new_values = {
                "Discovered Date": today,
                "Job Title": title,
                "Company": company,
                "Location": location or "Location not specified  -   verify",
                "Work Mode": mode,
                "Job Category": cat,
                "Experience Required": experience,
                "Fresher Eligible": experience_status,
                "Required Skills": skills_text,
                "Matched Skills": matched_text,
                "Match Score (%)": points,
                "Salary": salary,
                "Posted Date": posted,
                "Deadline": "",
                "Job URL": url,
                "Source": source_name,
                "Listing Status": listing_status,
                "Application Status": "Not applied",
                "Notes": assessment,
            }

            ws.append([
                new_values.get(ws.cell(1, i).value)
                for i in range(1, ws.max_column + 1)
            ])
            row_num = ws.max_row
            known[url] = row_num
            added += 1

    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    wb.close()

    print(
        f"\nListings retrieved: {len(alljobs)}\nNew rows added: {added}\nExisting rows updated: {updated}\nDuplicates/low-match/new unrelated listings skipped: {skipped}\nWorkbook: {path}"
    )
    print("Application statuses, discovery dates, deadlines, and existing notes are preserved.")
    print("Scores are keyword-based, not proof of eligibility. Verify each original listing.")
    print("Location and experience assessments are estimates; unclear eligibility requires verification.")


if __name__ == "__main__":
    main()
